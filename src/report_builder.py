"""Styled Excel workbook export (purple Lancer-Bot style, normal filters, no Excel Tables)."""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PURPLE, LIGHT_GRAY, WHITE, DARK, BORDER = "4B2E83", "F3F4F6", "FFFFFF", "202124", "B9A3D0"
PCT_COLS = {"Run_Rate", "Pass_Rate", "Rate", "Team_Baseline", "Lift"}


def _write_sheet(ws, title: str, data: pd.DataFrame):
    data = data if data is not None and not data.empty else pd.DataFrame({"Note": ["No rows for this section."]})
    ws.sheet_view.showGridLines = False
    ncols = max(1, len(data.columns))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws["A1"] = title
    ws["A1"].fill = PatternFill("solid", fgColor=PURPLE)
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center")
    border = Border(*(Side(style="thin", color=BORDER) for _ in range(4)))
    for c, name in enumerate(data.columns, 1):
        cell = ws.cell(3, c, str(name).replace("_", " "))
        cell.fill = PatternFill("solid", fgColor=PURPLE)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    pct_idx = {i for i, c in enumerate(data.columns, 1) if c in PCT_COLS}
    for r, row in enumerate(data.itertuples(index=False, name=None), 4):
        for c, v in enumerate(row, 1):
            if v is pd.NA or (isinstance(v, float) and pd.isna(v)):
                v = None
            elif hasattr(v, "item"):
                v = v.item()
            elif not isinstance(v, (int, float, str, bool)) and v is not None:
                v = str(v)
            cell = ws.cell(r, c, v)
            if c in pct_idx:
                cell.number_format = "0%"
            cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY if r % 2 == 0 else WHITE)
            cell.font = Font(name="Calibri", size=9, color=DARK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(ncols)}{max(3, ws.max_row)}"
    ws.row_dimensions[3].height = 30
    for c in range(1, ncols + 1):
        vals = [str(ws.cell(r, c).value or "") for r in range(3, min(ws.max_row, 103) + 1)]
        ws.column_dimensions[get_column_letter(c)].width = min(max(11, max(map(len, vals)) + 2), 40)


def build_workbook(sections: list[tuple[str, str, pd.DataFrame]]) -> bytes:
    """sections = [(sheet_name, title, dataframe), ...]. Sheet names are trimmed to Excel's 31 chars."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, title, df in sections:
        _write_sheet(wb.create_sheet(name[:31]), title, df)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
